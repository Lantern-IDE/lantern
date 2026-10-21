"""사람이 처음부터 쓰는 한국어 질문 표 (벤치마크의 한국어 질문은 LLM이 쓰고 LLM이 검토했으므로, 그 영향을 따로 재려고).

    python eval/bilingual/native_sheet.py export <작성표.xlsx> [--per-repo 4] [--seed 7]
    python eval/bilingual/native_sheet.py import <작성표.xlsx>    # 쓴 문항만 각 JSON의 q_ko_native로

개발 중에 보지 않은 문항(unseen)에서 저장소마다 같은 수를 무작위로 고른다 (기본 10개 저장소 × 4 = 40문항).
작성자는 영어 커밋 메시지만 보고, LLM이 쓴 한국어 질문·정답 파일·바뀐 코드는 보지 않는다.
측정: --lang ko-native (retrieval.mjs, bench/run.mjs, baselines.py). report.py가 같은 문항의 LLM 한국어 질문과 나란히 보여 준다.
openpyxl이 필요하다 (pip install openpyxl).
"""

import argparse
import glob
import json
import os
import random
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Arial"
HEAD = ["번호", "저장소", "영어 커밋 메시지", "한국어 질문 (직접 작성)", "작성자", "메모", "커밋 (뜻이 모호할 때만)"]
WIDTH = [16, 13, 60, 60, 10, 24, 16]
INPUT = PatternFill("solid", start_color="FFF6D5")
COL_KO, COL_WHO, COL_NOTE, COL_URL = 4, 5, 6, 7


def sets():
    for f in sorted(glob.glob(os.path.join(HERE, "*.json"))):
        with open(f, encoding="utf-8") as fh:
            yield f, json.load(fh)


def seen():
    out = set()
    for f in glob.glob(os.path.join(HERE, "..", "bench", "*.json")):
        with open(f, encoding="utf-8") as fh:
            out |= {q["commit"] for q in json.load(fh)["questions"]}
    return out


def export(path, per_repo, seed):
    wb = Workbook()
    guide = wb.active
    guide.title = "작성 안내"
    rows = [
        ["한국어 질문 직접 쓰는 법"],
        [],
        ["무엇", "오픈소스 저장소의 실제 커밋 메시지(영어)를 읽고, 그 일을 코딩 에이전트(Claude Code, Cursor 등)에게 시킨다면 한국어로 어떻게 말할지 씁니다."],
        ["", "번역이 아닙니다. 평소 내가 실제로 칠 말로 씁니다. 영어 낱말을 그대로 섞어도 되고, 짧아도 됩니다."],
        [],
        ["지킬 것 1", "번역기·ChatGPT·Claude 등 도구를 쓰지 않습니다. 영어 뜻을 몰라 사전을 찾는 것은 괜찮습니다."],
        ["지킬 것 2", "영어 메시지에 없는 정보(파일 이름, 함수 이름, 원인 추측)를 넣지 않습니다. 코드와 바뀐 내용(diff)은 보지 않습니다."],
        ["지킬 것 3", "코드 이름(MethodView, verify=False)은 보통 쓰는 대로 씁니다. 영어 그대로 써도, 한국어로 풀어도 됩니다."],
        ["지킬 것 4", "뜻이 모호하면 커밋 링크를 열어 메시지 본문까지만 읽습니다."],
        ["지킬 것 5", "도저히 질문으로 만들기 어려우면 비워 두고 메모에 이유를 씁니다."],
        [],
        ["예시 (벤치마크에 없는 가상의 커밋)"],
        HEAD[2:4],
        ["handle trailing slash in route matching", "라우트 매칭할 때 끝에 슬래시 붙어도 되게 해줘"],
        ["Fix crash when the config file is empty", "설정 파일 비어 있으면 죽는 거 고쳐줘"],
        [],
        ["노란 칸(한국어 질문, 작성자, 메모)만 채웁니다. 작성자 칸에는 이름이나 머리글자를 적습니다 (논문에는 사람 수만 밝힘)."],
        ["다 쓰면: python eval/bilingual/native_sheet.py import <이 파일>"],
    ]
    for r in rows:
        guide.append(r)
    for row in guide.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    guide["A1"].font = Font(name=FONT, size=14, bold=True)
    guide["A12"].font = Font(name=FONT, size=11, bold=True)
    for c in guide[13]:
        c.font = Font(name=FONT, size=11, bold=True)
    for col, w in zip("AB", [44, 90]):
        guide.column_dimensions[col].width = w

    ws = wb.create_sheet("질문")
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name=FONT, size=11, bold=True)
    skip = seen()
    rnd = random.Random(seed)
    n = 0
    for _, s in sets():
        cand = [(i, q) for i, q in enumerate(s["questions"], 1) if q["commit"] not in skip]
        rnd.shuffle(cand)
        for i, q in sorted(cand[:per_repo]):
            ws.append([f'{s["project"]}-{i:02d}', s["project"], q["q"], q.get("q_ko_native") or "", q.get("q_ko_native_by") or "", "", "열기"])
            r = ws.max_row
            ws.cell(r, COL_URL).hyperlink = f'{s["url"]}/commit/{q["commit"]}'
            for col in range(1, len(HEAD) + 1):
                ws.cell(r, col).font = Font(name=FONT, size=11)
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, COL_URL).font = Font(name=FONT, size=11, color="0563C1", underline="single")
            for col in (COL_KO, COL_WHO, COL_NOTE):
                ws.cell(r, col).fill = INPUT
            n += 1
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "D2"
    wb.active = 1
    wb.save(path)
    print(f"{path}: 질문 {n}개")


def import_(path):
    ws = load_workbook(path)["질문"]
    filled = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, _, en, ko, who = row[:5]
        if key:
            filled[key] = (en, str(ko or "").strip(), str(who or "").strip())
    written, writers = 0, set()
    for f, s in sets():
        for i, q in enumerate(s["questions"], 1):
            key = f'{s["project"]}-{i:02d}'
            if key not in filled:
                continue
            en, ko, who = filled[key]
            if en != q["q"]:
                sys.exit(f"{key}: 작성표의 영어 메시지가 JSON과 다릅니다. 작성표를 다시 만드세요 (export)")
            if ko:
                q["q_ko_native"] = ko
                q["q_ko_native_by"] = who or "?"
                written += 1
                writers.add(who or "?")
            else:
                q.pop("q_ko_native", None)
                q.pop("q_ko_native_by", None)
        with open(f, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(json.dumps(s, ensure_ascii=False, indent=2) + "\n")
    print(f"직접 쓴 질문 {written}/{len(filled)}개 · 작성자 {len(writers)}명")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument("cmd", choices=["export", "import"])
    ap.add_argument("path")
    ap.add_argument("--per-repo", type=int, default=4)
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args()
    if a.cmd == "export":
        export(a.path, a.per_repo, a.seed)
    else:
        import_(a.path)
