"""영·한 병렬 벤치마크의 한국어 질문 검토표.

    python eval/bilingual/ko_sheet.py export <검토표.xlsx>   # 질문 JSON → 검토표 (초안·이미 검토한 질문을 채움)
    python eval/bilingual/ko_sheet.py import <검토표.xlsx>   # 검토표 → 각 JSON의 q_ko (확인 칸에 표시한 문항만)

한국어 질문은 두 단계로 만든다.
- 초안(q_ko_draft): 이 엔진을 모르는 LLM에게 영어 질문과 작성 기준만 주고 쓰게 한 것
- 최종(q_ko): 저자가 초안을 모두 읽고 고친 것. 검토표의 '확인' 칸에 표시한 문항만 q_ko가 되고,
  표시하지 않은 문항은 q_ko가 비어 한국어 측정에서 빠진다
import는 초안을 고친 비율도 알려 준다 (논문에 밝힘). 검토표에는 정답 파일을 넣지 않는다.
openpyxl이 필요하다 (pip install openpyxl).
"""

import glob
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Arial"
HEAD = ["번호", "저장소", "영어 질문 (커밋 메시지)", "초안 (고치지 않음)", "한국어 질문 (고칠 것만 고침)", "확인", "메모", "커밋 (뜻이 모호할 때만)"]
WIDTH = [12, 13, 50, 50, 50, 7, 22, 16]
INPUT = PatternFill("solid", start_color="FFF6D5")
DRAFT = PatternFill("solid", start_color="F2F2F2")
COL_KO, COL_OK, COL_NOTE, COL_URL = 5, 6, 7, 8


def sets():
    for f in sorted(glob.glob(os.path.join(HERE, "*.json"))):
        with open(f, encoding="utf-8") as fh:
            yield f, json.load(fh)


def export(path):
    wb = Workbook()
    guide = wb.active
    guide.title = "검토 안내"
    rows = [
        ["한국어 질문 검토하는 법"],
        [],
        ["순서", "영어 질문을 읽고 → 초안을 읽고 → 어색하거나 틀렸으면 '한국어 질문' 칸을 고치고 → '확인' 칸에 O를 적습니다. 고칠 게 없어도 O를 적습니다."],
        ["", "O가 있는 문항만 벤치마크에 들어갑니다. 다 못 했으면 저장해서 넘겨도 됩니다 (O 친 것까지만 들어감)."],
        [],
        ["기준 1", "한국 개발자가 동료나 AI에게 실제로 묻듯이 자연스러운가. 번역투면 평소 쓰는 말로 고칩니다."],
        ["기준 2", "영어 질문에 없는 정보가 들어가지 않았나: 파일 이름, 함수 이름, 원인 추측 등. 있으면 뺍니다. 코드는 보지 않습니다."],
        ["기준 3", "뜻이 맞나: 부정(don't, without), 조건(only when) 등을 특히 봅니다."],
        ["기준 4", "코드 이름(MethodView, verify=False)과 고유명사는 영어 그대로 둬도 됩니다. 일반 낱말은 평소 쓰는 말로 (요청 본문 / 리퀘스트 바디 둘 다 괜찮음)."],
        ["기준 5", "말투는 자유입니다 (요청 / 질문 / 이슈 제목). 내가 실제로 이렇게 물을지를 기준으로 고칩니다."],
        ["기준 6", "뜻이 모호하면 커밋 링크를 열어 메시지 본문까지는 읽어도 됩니다. 바뀐 코드(diff)는 보지 않습니다."],
        ["기준 7", "도저히 질문으로 만들기 어려우면 '한국어 질문' 칸을 비우고 O를 적은 뒤 메모에 이유를 씁니다."],
        [],
        ["예시 (벤치마크에 없는 가상의 질문)"],
        HEAD[2:6],
        ["handle trailing slash in route matching", "route matching에서 trailing slash 처리하기", "라우트 매칭할 때 끝에 붙은 슬래시도 처리해줘", "O"],
        ["Fix crash when the config file is empty", "설정 파일이 비어 있을 때 크래시 나는 문제", "설정 파일이 비어 있을 때 크래시 나는 문제", "O"],
        [],
        ["노란 칸(한국어 질문, 확인, 메모)만 고칩니다. 다 하면 python eval/bilingual/ko_sheet.py import <이 파일>로 벤치마크에 넣습니다."],
    ]
    for r in rows:
        guide.append(r)
    for row in guide.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    guide["A1"].font = Font(name=FONT, size=14, bold=True)
    guide["A14"].font = Font(name=FONT, size=11, bold=True)
    for c in guide[15]:
        c.font = Font(name=FONT, size=11, bold=True)
    for col, w in zip("ABCD", [44, 90, 44, 8]):
        guide.column_dimensions[col].width = w

    ws = wb.create_sheet("질문")
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name=FONT, size=11, bold=True)
    n = 0
    for _, s in sets():
        for i, q in enumerate(s["questions"], 1):
            reviewed = bool(q.get("q_ko_reviewed"))
            ko = (q.get("q_ko") or "") if reviewed else (q.get("q_ko_draft") or "")
            ws.append([f'{s["project"]}-{i:02d}', s["project"], q["q"], q.get("q_ko_draft") or "", ko, "O" if reviewed else "", q.get("q_ko_note") or "", "열기"])
            r = ws.max_row
            ws.cell(r, COL_URL).hyperlink = f'{s["url"]}/commit/{q["commit"]}'
            for col in range(1, len(HEAD) + 1):
                ws.cell(r, col).font = Font(name=FONT, size=11)
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, COL_URL).font = Font(name=FONT, size=11, color="0563C1", underline="single")
            ws.cell(r, 4).fill = DRAFT
            for col in (COL_KO, COL_OK, COL_NOTE):
                ws.cell(r, col).fill = INPUT
            ws.cell(r, COL_OK).alignment = Alignment(horizontal="center", vertical="top")
            n += 1
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{chr(64 + len(HEAD))}{ws.max_row}"
    wb.active = 1
    wb.save(path)
    print(f"{path}: 질문 {n}개")


def import_(path):
    ws = load_workbook(path)["질문"]
    filled = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, _, en, _draft, ko, ok, note = row[:7]
        if key:
            filled[key] = (en, str(ko or "").strip(), bool(str(ok or "").strip()), str(note or "").strip())
    reviewed = changed = empty = total = 0
    for f, s in sets():
        for i, q in enumerate(s["questions"], 1):
            total += 1
            key = f'{s["project"]}-{i:02d}'
            if key not in filled:
                continue
            en, ko, ok, note = filled[key]
            if en != q["q"]:
                sys.exit(f"{key}: 검토표의 영어 질문이 JSON과 다릅니다. 검토표를 다시 만드세요 (export)")
            if ok:
                q["q_ko"] = ko or None
                q["q_ko_reviewed"] = True
                reviewed += 1
                changed += bool(ko) and ko != q.get("q_ko_draft")
                empty += not ko
            else:
                q["q_ko"] = None
                q.pop("q_ko_reviewed", None)
            if note:
                q["q_ko_note"] = note
            else:
                q.pop("q_ko_note", None)
        with open(f, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(json.dumps(s, ensure_ascii=False, indent=2) + "\n")
    rate = f" ({changed / reviewed * 100:.0f}%)" if reviewed else ""
    print(f"검토 {reviewed}/{total}개 · 초안을 고친 것 {changed}개{rate} · 비운 것 {empty}개")


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("export", "import"):
        sys.exit(__doc__)
    (export if sys.argv[1] == "export" else import_)(sys.argv[2])
