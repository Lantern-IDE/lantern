"""한국어 모국어 개발자가 LLM이 쓴 한국어 질문을 평가만 하는 표. 질문은 고치지 않는다 (이미 잰 질문이므로).
평가자 A(저자)와 B(결과를 보지 않은 다른 개발자)가 같은 문항을 따로 평가하고, 일치도를 낸다.

    python eval/bilingual/ko_rating.py export <평가표.xlsx> --rater A|B [--n 60] [--seed 11]
    python eval/bilingual/ko_rating.py import <평가표.xlsx> --rater A|B   # ratings/ko-ratings-<평가자>.json에 쓰고 요약 출력
    python eval/bilingual/ko_rating.py agree                              # 두 평가자의 일치도 (같이 평가한 문항)

기준은 평가 전에 정해 둔다: ko-rating-plan.md. 개발 중에 보지 않은 문항에서 무작위로 고른다.
같은 시드면 처음 30문항은 예전 30문항 표본과 같다 (random.sample은 앞에서부터 같은 순서로 뽑음).
openpyxl이 필요하다 (pip install openpyxl).
"""

import argparse
import glob
import json
import os
import random
import sys
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
# 하위 폴더에: eval/bilingual/*.json은 모든 측정 스크립트가 문항 세트로 읽는다
OUT_DIR = os.path.join(HERE, "ratings")
FONT = "Arial"
HEAD = ["번호", "저장소", "영어 커밋 제목", "LLM이 쓴 한국어 질문", "자연스러움 (1~3)", "뜻", "메모"]
WIDTH = [16, 13, 55, 55, 14, 12, 30]
INPUT = PatternFill("solid", start_color="FFF6D5")
NATURAL = [1, 2, 3]
MEANING = ["같음", "일부", "다름"]


def guide(rater):
    return [
        [f"LLM 한국어 질문 평가 — 평가자 {rater} (질문은 고치지 않습니다)"],
        [],
        ["자연스러움", "3 = 한국 개발자가 실제로 이렇게 묻겠다 / 2 = 뜻은 통하지만 번역투나 어색한 데가 있다 / 1 = 이렇게는 묻지 않겠다"],
        ["뜻", "같음 = 영어 제목과 뜻이 같다 / 일부 = 조건·부정·대상 일부가 빠지거나 달라졌다 / 다름 = 뜻이 틀렸다"],
        ["메모", "어색한 곳이나 틀린 곳을 짧게 (선택)"],
        [],
        ["지킬 것", "질문을 고치지 않습니다. 정답 파일이나 검색 결과는 보지 않습니다. 다른 평가자의 표는 보지 않습니다. 한 번 적은 평가는 바꾸지 않습니다."],
        ["기준", "평가 기준은 eval/bilingual/ko-rating-plan.md에 미리 적어 두었습니다. 평가자 A(저자)와 B(다른 한국어 모국어 개발자)가 같은 문항을 따로 평가합니다."],
        [],
        [f"다 하면: python eval/bilingual/ko_rating.py import <이 파일> --rater {rater}"],
    ]


def sets():
    for f in sorted(glob.glob(os.path.join(HERE, "*.json"))):
        with open(f, encoding="utf-8") as fh:
            yield f, json.load(fh)


def export(path, n, seed, rater):
    seen = set()
    for f in glob.glob(os.path.join(HERE, "..", "bench", "*.json")):
        with open(f, encoding="utf-8") as fh:
            seen |= {q["commit"] for q in json.load(fh)["questions"]}
    pool = [(s["project"], i, q) for _, s in sets() for i, q in enumerate(s["questions"], 1) if q["commit"] not in seen and q.get("q_ko")]
    picked = random.Random(seed).sample(pool, n)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "평가 기준"
    for r in guide(rater):
        sheet.append(r)
    for row in sheet.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    for col, w in zip("AB", [14, 100]):
        sheet.column_dimensions[col].width = w
    ws = wb.create_sheet("평가")
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name=FONT, size=11, bold=True)
    nat = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    mean = DataValidation(type="list", formula1=f'"{",".join(MEANING)}"', allow_blank=True)
    ws.add_data_validation(nat)
    ws.add_data_validation(mean)
    for proj, i, q in picked:
        ws.append([f"{proj}-{i:02d}", proj, q["q"], q["q_ko"], "", "", ""])
        r = ws.max_row
        for col in range(1, len(HEAD) + 1):
            ws.cell(r, col).font = Font(name=FONT, size=11)
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (5, 6, 7):
            ws.cell(r, col).fill = INPUT
        nat.add(ws.cell(r, 5))
        mean.add(ws.cell(r, 6))
    for k, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + k)].width = w
    ws.freeze_panes = "E2"
    wb.active = 1
    wb.save(path)
    print(f"{path}: 평가자 {rater}, {n}문항")


def out_path(rater):
    return os.path.join(OUT_DIR, f"ko-ratings-{rater}.json")


def import_(path, rater):
    ws = load_workbook(path)["평가"]
    ko = {f'{s["project"]}-{i:02d}': q["q_ko"] for _, s in sets() for i, q in enumerate(s["questions"], 1)}
    rows = []
    for key, proj, en, q_ko, nat, mean, note in ws.iter_rows(min_row=2, values_only=True):
        if not key:
            continue
        if ko.get(key) != q_ko:
            sys.exit(f"{key}: 평가표의 한국어 질문이 JSON과 다릅니다 (질문은 고치지 않습니다)")
        if nat is None and not mean:
            continue
        if (nat is not None and int(nat) not in NATURAL) or (mean and mean not in MEANING):
            sys.exit(f"{key}: 기준에 없는 값 ({nat}, {mean})")
        rows.append({"id": key, "naturalness": int(nat) if nat is not None else None, "meaning": mean or None, "note": note or ""})
    os.makedirs(OUT_DIR, exist_ok=True)
    who = "author (native Korean speaker and developer)" if rater == "A" else "second native Korean developer, blind to results"
    with open(out_path(rater), "w", encoding="utf-8", newline="\n") as fh:
        json.dump({"rater": rater, "who": who, "plan": "ko-rating-plan.md", "ratings": rows}, fh, ensure_ascii=False, indent=1)
    nat = Counter(r["naturalness"] for r in rows if r["naturalness"])
    mean = Counter(r["meaning"] for r in rows if r["meaning"])
    print(f"평가자 {rater}: {len(rows)}문항 → {out_path(rater)}")
    print(f"자연스러움: 3 {nat[3]} · 2 {nat[2]} · 1 {nat[1]}")
    print(f"뜻: 같음 {mean['같음']} · 일부 {mean['일부']} · 다름 {mean['다름']}")


def kappa(a, b, cats, quadratic=False):
    """Cohen의 카파. quadratic이면 순서 범주의 2차 가중 카파"""
    k, n = len(cats), len(a)
    idx = {c: i for i, c in enumerate(cats)}
    obs = [[0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[idx[x]][idx[y]] += 1
    pa = [sum(obs[i]) / n for i in range(k)]
    pb = [sum(obs[i][j] for i in range(k)) / n for j in range(k)]
    w = [[((i - j) / (k - 1)) ** 2 if quadratic else float(i != j) for j in range(k)] for i in range(k)]
    seen = sum(w[i][j] * obs[i][j] / n for i in range(k) for j in range(k))
    chance = sum(w[i][j] * pa[i] * pb[j] for i in range(k) for j in range(k))
    return 1 - seen / chance if chance else float("nan")


def agree():
    got = {}
    for r in ("A", "B"):
        if not os.path.exists(out_path(r)):
            sys.exit(f"{out_path(r)} 없음: 두 평가자 모두 import한 뒤 실행")
        with open(out_path(r), encoding="utf-8") as fh:
            got[r] = {x["id"]: x for x in json.load(fh)["ratings"]}
    for crit, cats, quad in (("naturalness", NATURAL, True), ("meaning", MEANING, False)):
        ids = [i for i in got["A"] if i in got["B"] and got["A"][i][crit] is not None and got["B"][i][crit] is not None]
        a, b = [got["A"][i][crit] for i in ids], [got["B"][i][crit] for i in ids]
        if not ids:
            print(f"{crit}: 같이 평가한 문항 없음")
            continue
        same = sum(x == y for x, y in zip(a, b)) / len(ids) * 100
        name = "2차 가중 카파" if quad else "Cohen 카파"
        print(f"{crit}: {len(ids)}문항, 일치 {same:.1f}%, {name} {kappa(a, b, cats, quad):.2f}")
        for r, v in (("A", a), ("B", b)):
            c = Counter(v)
            print(f"  {r}: " + " · ".join(f"{x} {c[x]}" for x in cats))
    differ = sorted(i for r in ("A", "B") for i, x in got[r].items() if x["meaning"] == "다름")
    print("뜻이 '다름'인 문항 (어느 평가자든):", ", ".join(sorted(set(differ))) or "없음")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument("cmd", choices=["export", "import", "agree"])
    ap.add_argument("path", nargs="?")
    ap.add_argument("--rater", choices=["A", "B"], default="A")
    ap.add_argument("--n", type=int, default=60)
    ap.add_argument("--seed", type=int, default=11)
    a = ap.parse_args()
    if a.cmd != "agree" and not a.path:
        ap.error("export·import에는 평가표 경로가 필요합니다")
    {"export": lambda: export(a.path, a.n, a.seed, a.rater), "import": lambda: import_(a.path, a.rater), "agree": agree}[a.cmd]()
